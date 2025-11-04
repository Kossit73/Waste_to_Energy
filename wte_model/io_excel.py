"""Excel IO helpers for the waste-to-energy model."""
from __future__ import annotations

from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from .inputs import (
    Timeline,
    TechAssumptions,
    RevenueAssumptions,
    CostAssumptions,
    FinanceAssumptions,
    WTEMasterInputs,
)

NAME_MAP: Dict[str, List[str]] = {
    "project_start_date": ["project_start_date", "proj_start_date", "start_date"],
    "pre_op_months": ["pre_op_months", "preop_months", "pre_operations_months"],
    "dev_months": ["dev_months", "development_months"],
    "build_months": ["build_months", "construction_months", "c_months"],
    "cod_date": ["COD_date", "cod_date", "date_COD"],
    "eoc_date": ["eoc_date", "EOC_date", "end_of_construction_date"],
    "periods_per_year": ["periods_per_year", "ppy", "model_ppy"],
    "msw_tonnes_pa": ["msw_tonnes_pa", "waste_tonnage_pa", "msw_annual_tonnage", "annual_waste_tonnes"],
    "lhv_mj_per_kg": ["LHV_MJ_per_kg", "lhv_mj_kg", "LHV", "lhv"],
    "boiler_efficiency": ["boiler_efficiency", "boiler_eta", "steam_cycle_efficiency"],
    "electrical_efficiency": ["electrical_efficiency", "net_electrical_efficiency", "elec_eff_net", "eta_electric_net"],
    "availability": ["availability", "availability_pa", "on_stream_factor", "uptime_factor"],
    "parasitic_load_frac": ["parasitic_load_frac", "parasitic_load", "aux_power_frac", "house_load_frac"],
    "plant_capacity_MW": ["plant_capacity_MW", "net_export_capacity_MW", "net_capacity_MW"],
    "power_export_MWh_pa": ["power_export_MWh_pa", "annual_export_MWh"],
    "ppa_price_usd_per_mwh": ["ppa_tariff_USD_per_MWh", "ppa_price", "ppa_tariff", "tariff_usd_mwh"],
    "ppa_escalation": ["ppa_escalation", "ppa_cpi", "ppa_escalation_pa"],
    "gate_fee_usd_per_t": ["gate_fee_USD_per_ton", "tipping_fee", "gate_fee", "gate_fee_usd_t"],
    "gate_fee_escalation": ["gate_fee_escalation", "tipping_fee_escalation", "gate_cpi"],
    "heat_price_usd_per_mwh": ["heat_price_usd_per_mwh", "heat_tariff", "thermal_tariff"],
    "other_metal_usd_per_t": ["metal_recovery_usd_per_t", "metals_revenue_usd_t"],
    "other_ash_usd_per_t": ["ash_revenue_usd_per_t", "ash_sales_usd_t"],
    "other_escalation": ["other_escalation", "byproduct_escalation"],
    "capex_total_usd": ["capex_total_usd", "CAPEX_total", "total_capex", "capex_total"],
    "capex_spend_profile": ["capex_spend_profile", "capex_profile", "capex_draw_profile", "capex_spread"],
    "fixed_om_usd_pa": ["fixed_om_usd_pa", "fixed_om_pa", "fixed_opex_pa"],
    "variable_om_usd_per_t": ["variable_om_usd_per_t", "variable_om_t", "var_opex_usd_t"],
    "landfill_disposal_usd_per_t": ["landfill_disposal_usd_per_t", "ash_disposal_usd_t", "residue_disposal_usd_t"],
    "insurance_pct_of_capex_pa": ["insurance_pct_of_capex_pa", "insurance_pct_capex", "insurance_pct"],
    "maintenance_pct_of_capex_pa": ["maintenance_pct_of_capex_pa", "maintenance_pct_capex", "maintenance_pct"],
    "opex_escalation": ["opex_escalation", "opex_cpi", "om_escalation"],
    "debt_ratio": ["debt_ratio", "gearing", "debt_to_capital"],
    "interest_rate": ["debt_interest_rate", "interest_rate", "loan_interest_rate"],
    "tenor_years": ["tenor_years", "debt_tenor_years", "loan_tenor_years"],
    "grace_years": ["grace_years", "grace_period_years"],
    "upfront_fee_pct": ["upfront_fee_pct", "arrangement_fee_pct", "debt_fee_pct"],
    "dscr_min": ["dscr_min", "min_dscr", "covenant_dscr_min"],
    "tax_rate": ["tax_rate", "corp_tax_rate"],
    "depr_years": ["depr_years", "depreciation_years"],
    "working_cap_days": ["working_cap_days", "wc_days"],
    "discount_rate": ["discount_rate", "wacc", "hurdle_rate"],
}

PERCENT_KEYS = {
    "boiler_efficiency",
    "electrical_efficiency",
    "availability",
    "parasitic_load_frac",
    "ppa_escalation",
    "gate_fee_escalation",
    "other_escalation",
    "opex_escalation",
    "debt_ratio",
    "interest_rate",
    "upfront_fee_pct",
    "tax_rate",
    "insurance_pct_of_capex_pa",
    "maintenance_pct_of_capex_pa",
}


def _maybe_percent(key: str, val: Any) -> Any:
    try:
        x = float(val)
    except Exception:
        return val
    if key in PERCENT_KEYS:
        if 0.0 <= x <= 1.0:
            return x
        if 1.0 < x <= 100.0:
            return x / 100.0
    return x


def _get_defined_single(wb: Workbook, name: str) -> Optional[Any]:
    dn = wb.defined_names.get(name)
    if not dn:
        return None
    dests = list(dn.destinations)
    if len(dests) != 1:
        if len(dests) == 0:
            return None
        sheet, ref = dests[0]
        cell = wb[sheet][ref]
        return cell.value
    sheet, ref = dests[0]
    ws = wb[sheet]
    if ":" in ref:
        start = ref.split(":")[0]
        return ws[start].value
    return ws[ref].value


def _get_defined_range(wb: Workbook, name: str) -> Optional[List[float]]:
    dn = wb.defined_names.get(name)
    if not dn:
        return None
    dests = list(dn.destinations)
    if len(dests) == 0:
        return None
    sheet, ref = dests[0]
    ws = wb[sheet]
    values: List[float] = []
    try:
        if ":" in ref:
            for row in ws[ref]:
                for cell in row:
                    if cell.value not in (None, ""):
                        try:
                            values.append(float(cell.value))
                        except Exception:
                            pass
        else:
            val = ws[ref].value
            if val is not None:
                values.append(float(val))
    except Exception:
        return None
    return values or None


def _get_named_first(wb: Workbook, aliases: List[str]) -> Optional[Any]:
    for alias in aliases:
        value = _get_defined_single(wb, alias)
        if value is not None:
            return value
    return None


def _get_named_range_first(wb: Workbook, aliases: List[str]) -> Optional[List[float]]:
    for alias in aliases:
        value = _get_defined_range(wb, alias)
        if value:
            return value
    return None


def _safe_float(value: Any, default: float) -> float:
    try:
        return float(value)
    except Exception:
        return default


def load_inputs_from_xlsm(path: str) -> WTEMasterInputs:
    wb = load_workbook(path, read_only=False, data_only=True, keep_vba=True)

    singles = {key: _get_named_first(wb, aliases) for key, aliases in NAME_MAP.items()}
    capex_profile = _get_named_range_first(wb, NAME_MAP["capex_spend_profile"])

    build_months = int(_safe_float(singles.get("build_months"), 24))
    ppy = int(_safe_float(singles.get("periods_per_year"), 4))
    timeline = Timeline(years=25, build_months=build_months, start_year=2026, periods_per_year=ppy)

    tech = TechAssumptions(
        msw_tonnes_pa=_safe_float(singles.get("msw_tonnes_pa"), 300_000.0),
        lhv_mj_per_kg=_safe_float(singles.get("lhv_mj_per_kg"), 9.0),
        boiler_efficiency=_maybe_percent("boiler_efficiency", singles.get("boiler_efficiency") or 0.85),
        electrical_efficiency=_maybe_percent(
            "electrical_efficiency", singles.get("electrical_efficiency") or 0.22
        ),
        availability=_maybe_percent("availability", singles.get("availability") or 0.90),
        parasitic_load_frac=_maybe_percent("parasitic_load_frac", singles.get("parasitic_load_frac") or 0.10),
    )

    other_per_t = 0.0
    if singles.get("other_metal_usd_per_t") is not None:
        other_per_t += _safe_float(singles["other_metal_usd_per_t"], 0.0)
    if singles.get("other_ash_usd_per_t") is not None:
        other_per_t += _safe_float(singles["other_ash_usd_per_t"], 0.0)

    revenue = RevenueAssumptions(
        ppa_price_usd_per_mwh=_safe_float(singles.get("ppa_price_usd_per_mwh"), 110.0),
        gate_fee_usd_per_t=_safe_float(singles.get("gate_fee_usd_per_t"), 25.0),
        heat_price_usd_per_mwh=_safe_float(singles.get("heat_price_usd_per_mwh"), 0.0),
        metal_recovery_usd_per_t=_safe_float(singles.get("other_metal_usd_per_t"), 0.0),
        ash_revenue_usd_per_t=_safe_float(singles.get("other_ash_usd_per_t"), 0.0),
        ppa_escalation=_maybe_percent("ppa_escalation", singles.get("ppa_escalation") or 0.0),
        gate_fee_escalation=_maybe_percent("gate_fee_escalation", singles.get("gate_fee_escalation") or 0.0),
        other_escalation=_maybe_percent("other_escalation", singles.get("other_escalation") or 0.0),
    )

    costs = CostAssumptions(
        capex_total_usd=_safe_float(singles.get("capex_total_usd"), 180_000_000.0),
        capex_spend_profile=capex_profile,
        fixed_om_usd_pa=_safe_float(singles.get("fixed_om_usd_pa"), 10_000_000.0),
        variable_om_usd_per_t=_safe_float(singles.get("variable_om_usd_per_t"), 15.0),
        landfill_disposal_usd_per_t=_safe_float(singles.get("landfill_disposal_usd_per_t"), 5.0),
        insurance_pct_of_capex_pa=_maybe_percent(
            "insurance_pct_of_capex_pa", singles.get("insurance_pct_of_capex_pa") or 0.0075
        ),
        maintenance_pct_of_capex_pa=_maybe_percent(
            "maintenance_pct_of_capex_pa", singles.get("maintenance_pct_of_capex_pa") or 0.02
        ),
        opex_escalation=_maybe_percent("opex_escalation", singles.get("opex_escalation") or 0.02),
    )

    finance = FinanceAssumptions(
        debt_ratio=_maybe_percent("debt_ratio", singles.get("debt_ratio") or 0.70),
        interest_rate=_maybe_percent("interest_rate", singles.get("interest_rate") or 0.08),
        tenor_years=int(_safe_float(singles.get("tenor_years"), 12)),
        grace_years=int(_safe_float(singles.get("grace_years"), 2)),
        upfront_fee_pct=_maybe_percent("upfront_fee_pct", singles.get("upfront_fee_pct") or 0.01),
        dscr_min=_safe_float(singles.get("dscr_min"), 1.20),
        tax_rate=_maybe_percent("tax_rate", singles.get("tax_rate") or 0.25),
        depr_years=int(_safe_float(singles.get("depr_years"), 15)),
        working_cap_days=int(_safe_float(singles.get("working_cap_days"), 30)),
        discount_rate=_safe_float(singles.get("discount_rate"), 0.10),
    )

    mapped_debug = {
        "timeline": {"build_months": build_months, "periods_per_year": ppy},
        "tech": tech.__dict__,
        "revenue": {
            "ppa_price_usd_per_mwh": revenue.ppa_price_usd_per_mwh,
            "gate_fee_usd_per_t": revenue.gate_fee_usd_per_t,
            "heat_price_usd_per_mwh": revenue.heat_price_usd_per_mwh,
            "metal_recovery_usd_per_t": revenue.metal_recovery_usd_per_t,
            "ash_revenue_usd_per_t": revenue.ash_revenue_usd_per_t,
            "ppa_escalation": revenue.ppa_escalation,
            "gate_fee_escalation": revenue.gate_fee_escalation,
            "other_escalation": revenue.other_escalation,
        },
        "costs": {
            "capex_total_usd": costs.capex_total_usd,
            "capex_spend_profile": capex_profile[:10] if capex_profile else None,
            "fixed_om_usd_pa": costs.fixed_om_usd_pa,
            "variable_om_usd_per_t": costs.variable_om_usd_per_t,
            "landfill_disposal_usd_per_t": costs.landfill_disposal_usd_per_t,
            "insurance_pct_of_capex_pa": costs.insurance_pct_of_capex_pa,
            "maintenance_pct_of_capex_pa": costs.maintenance_pct_of_capex_pa,
            "opex_escalation": costs.opex_escalation,
        },
        "finance": finance.__dict__,
    }
    WTEMasterInputs._MAPPED_DEBUG = mapped_debug  # type: ignore[attr-defined]

    return WTEMasterInputs(
        timeline=timeline,
        tech=tech,
        revenue=revenue,
        costs=costs,
        finance=finance,
    )
